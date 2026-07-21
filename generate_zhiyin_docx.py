import os
import subprocess
import sys

def install_and_import(package):
    try:
        import docx
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
        import docx

install_and_import('python-docx')
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create Document
doc = Document()

# Title
title = doc.add_heading('智音(ZhiYin) HBG-V1.0 全息交互玩具功能介绍书', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 1. 项目简介
doc.add_heading('一、项目简介', level=1)
p = doc.add_paragraph('智音 HBG-V1.0 是一款基于 ESP32-S3 与全息光学显示的多模态 AI 交互硬件。系统将大语言模型与硬件控制深度结合，提供实时语音智能陪伴与可视化的动态角色互动。')

# 2. 核心功能
doc.add_heading('二、核心功能介绍', level=1)

func1 = doc.add_paragraph()
func1.add_run('1. 实时 AI 语音对话\n').bold = True
func1.add_run('取消传统语言唤醒，利用电容触摸即刻发起对话。通过直连云端大语言模型与高效音频流解码，实现极低延迟的自然语音交互。')

func2 = doc.add_paragraph()
func2.add_run('2. 专属动态角色视觉交互\n').bold = True
func2.add_run('搭载高清 LCD 屏与动态渲染引擎，屏幕呈现专属宠物角色。角色拥有完整状态机，可根据实时对话情境做出喜悦、疑惑、播报等同步动作反馈。')

func3 = doc.add_paragraph()
func3.add_run('3. 视听同步实时字幕\n').bold = True
func3.add_run('针对复杂声学环境，设备在语音播报的同时，自动捕获下行数据流并在屏幕同步展示文本字幕，实现精准的看听同步。')

func4 = doc.add_paragraph()
func4.add_run('4. 大模型智能硬件控制\n').bold = True
func4.add_run('具备 AI 直接调度硬件能力。用户通过自然语言下达指令（如调节音量、控制灯光），云端解析语义后下发特定报文至单片机完成物理控制。')

func5 = doc.add_paragraph()
func5.add_run('5. 网络流媒体音频播放\n').bold = True
func5.add_run('系统内嵌流媒体解析模块，具备桌面网络音箱功能。用户可按需点播在线音频流，提供日常背景音乐陪伴。')

# 3. 购物清单/物料表
doc.add_heading('三、物料采购清单 (BOM)', level=1)
p_bom = doc.add_paragraph('以下为项目外围物料与结构件清单（淘宝零售参考价）：')

table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '物料名称'
hdr_cells[1].text = '单价 (元 / 人民币)'

records = [
    ('FPC-10P延长板0.5MM转1.0MM间距 焊好翻盖下接', '1.50'),
    ('WS2812幻彩RGB可编程LED防水灯带', '5.00'),
    ('TTP223触摸模块传感器电容式点动', '0.50'),
    ('电脑AR增透膜显示器高清防反光', '15.00'),
    ('304十字圆头自攻螺丝 100个', '3.00'),
    ('腔体喇叭扬声器，箱体喇叭全系列', '8.00'),
    ('金逸晨2.8英寸TFT液晶屏 ST7789小屏幕240x320显', '25.00'),
    ('立式带螺纹音视频耳机插座', '0.50'),
    ('软排线AWM20624 FFC/FPC软排线', '1.00'),
    ('泡泡机开关按钮配件开关diy轻触', '0.50'),
    ('定制亚克力板', '35.00'),
    ('3D建模', '45.00'),
    ('ESP32S3主板', '56.42'),
]

total = 0
for name, price in records:
    row_cells = table.add_row().cells
    row_cells[0].text = name
    row_cells[1].text = f"¥ {price}"
    total += float(price)

total_row = table.add_row().cells
total_row[0].text = '总计预估'
total_row[1].text = f"¥ {total:.2f}"

output_path = r'a:\Study\esp32\Project\ZhiYin-HBG-V1.0\智音全息玩具功能介绍说明书.docx'
doc.save(output_path)
print(f"Document updated successfully at {output_path}")
