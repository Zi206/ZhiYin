import os
import subprocess
import sys

def install_and_import(package):
    try:
        import docx
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
        import docx

install_and_import('python-docx')
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_heading('智音项目采购报销明细表', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('以下为项目相关采购及链接清单，请核对：')

# Create Table
table = doc.add_table(rows=1, cols=5)
table.style = 'Table Grid'

# Header
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '商品名称'
hdr_cells[1].text = '购买链接'
hdr_cells[2].text = '数量'
hdr_cells[3].text = '发票能否开具'
hdr_cells[4].text = '备注'

# Data
data = [
    (
        'FFC/FPC软排线0.5mm扁平连接线反/同向5cm-50cm 4p6p8p12p40p-80p',
        'https://e.tb.cn/h.RuxEpi6CaLvr5O5?tk=Ko22gk1geil',
        '1个',
        '能',
        '10P-0.05米-同向-10条'
    ),
    (
        'PC耐力板透明5mm3mm2mm阳光板车雨棚采光屋顶遮阳聚碳酸酯板定制',
        'https://e.tb.cn/h.REfe5Sj4XVG9yNB?tk=Tfjwgk1gMRu',
        '1个',
        '能',
        '定制透明罩子'
    ),
    (
        'TTP223 224 226触摸传感器触摸按键感应模块电容式点动型接近开关',
        'https://e.tb.cn/h.RubBKXXhdxcacwg?tk=NaE3gk16UF8',
        '1个',
        '能',
        '电容式触摸开关按键模块 红色'
    ),
    (
        '触摸按键开关感应模块 3V-30V点动/锁存 双稳态轻触开关 LED灯带',
        'https://e.tb.cn/h.RFEoHp0WuYcZ3T2?tk=myYYgk1TGmA',
        '1个',
        '能',
        '触摸按键开关感应模块 3线 蓝灯'
    ),
    (
        '【专业好打 适用拓竹】eSUN易生pla耗材3d打印耗材料pla basic哑光matte丝绸silk透明PLA+快速3D打印机耗材料',
        'https://e.tb.cn/h.REf4oqquOkwz9Eq?tk=fmtegk1hMgb',
        '1个',
        '能',
        'PLA Basic无卷盘 粉色 1KG'
    ),
    (
        '快来捡漏【软件资源分享，Claude账号直供】',
        'https://m.tb.cn/h.RFEKejx?tk=9Ch2gk17Ix4',
        '1个',
        '否',
        'Claude max20x成品号（非成品号没有人脸认证和接码，容易封号）'
    ),
    (
        'Claude max 20x成品号',
        '',
        '1个',
        '能',
        '需要另加微信: kaixidawang  Claude max 20x成品号'
    ),
    (
        '定制玻璃模型',
        '',
        '1个',
        '能',
        '需要提供3D打印文件给他'
    )
]

# Populate Table
for name, link, count, invoice, remark in data:
    row_cells = table.add_row().cells
    row_cells[0].text = name
    row_cells[1].text = link
    row_cells[2].text = count
    row_cells[3].text = invoice
    row_cells[4].text = remark

# Adjust column widths manually
for row in table.rows:
    row.cells[0].width = Inches(2.5)
    row.cells[1].width = Inches(2.0)
    row.cells[2].width = Inches(0.8)
    row.cells[3].width = Inches(1.0)
    row.cells[4].width = Inches(2.0)

output_path = r'a:\Study\esp32\Project\ZhiYin-HBG-V1.0\智音项目采购报销明细表.docx'
doc.save(output_path)
print(f"Reimbursement document generated at {output_path}")
