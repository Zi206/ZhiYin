import docx
from docx.shared import Pt

doc = docx.Document('02-嵌入式固件架构设计说明书.docx')

# 找到包含架构图的段落，已知在 '1.1 分层架构' 之后
target_idx = -1
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == '1.1 分层架构':
        target_idx = i + 1
        break

if target_idx != -1 and target_idx < len(doc.paragraphs):
    p = doc.paragraphs[target_idx]
    
    # 清空该段落的内容(也就是把原来的图删掉)
    p.clear()
    
    text_content = """【Application Layer】
音频状态机 (FSM) | UI 管理器 (LVGL) | MCP 工具注册与调度 (JSON-RPC 2.0)

【Protocol Layer】
WebSocket (TLS1.3) | MQTT+UDP (TLS1.2) | BinaryProtocol2 (Opus Frame封装)

【Audio Pipeline】
AEC3 回声消除 → NS/AGC 降噪增益 → VAD 端点检测 → Opus 编解码 → DMA I2S 传输

【Hardware Abstraction】
主控BSP | 协处理器BSP | LCD专用驱动 | Codec 驱动 | IMU/触摸传感器"""

    # 插入新的纯文本，并设置格式
    run = p.add_run(text_content)
    run.font.size = Pt(10)
    run.font.name = 'Consolas'

    # 保存
    doc.save('02-嵌入式固件架构设计说明书.docx')
    print("Replace success")
else:
    print("Cannot find target paragraph")
